import os
import re
import logging
from typing import Any, List, Dict, Optional
from pathlib import Path
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)


class DocumentParser:
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.pptx', '.txt', '.html', '.md'}

    def __init__(self, encoding: Optional[str] = 'utf-8'):
        self.chunk_size = 500
        self.chunk_overlap = 50
        self.encoding = encoding

    def parse_file(self, file_path: str) -> List[Dict[str, Any]]:
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            return []

        parser_method = getattr(self, f'_parse_{ext[1:]}', None)
        if parser_method:
            return parser_method(file_path)
        return []

    def _parse_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            with pdfplumber.open(file_path) as pdf:
                page_texts: list[tuple[int, str]] = []
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    page_texts.append((page_num, f"\n\n--- Page {page_num} ---\n\n" + text))

                # 逐页切，避免把整本 PDF 拼成大字符串
                for page_num, page_text in page_texts:
                    text_chunks = self._split_into_chunks(page_text)
                    for chunk in text_chunks:
                        chunks.append({
                            "content": chunk,
                            "source": os.path.basename(file_path),
                            "page": page_num,  # 真实 PDF 页码（而非 chunk 序号）
                            "doc_type": "pdf"
                        })
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}")
        return chunks

    def _parse_docx(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            doc = Document(file_path)
            full_text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

            text_chunks = self._split_into_chunks(full_text)
            for i, chunk in enumerate(text_chunks):
                chunks.append({
                    "content": chunk,
                    "source": os.path.basename(file_path),
                    "page": i + 1,
                    "doc_type": "docx"
                })
        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {e}")
        return chunks

    def _parse_xlsx(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            wb = load_workbook(file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_text = f"=== Sheet: {sheet_name} ===\n"

                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    if row_text.strip():
                        sheet_text += row_text + "\n"

                text_chunks = self._split_into_chunks(sheet_text)
                for i, chunk in enumerate(text_chunks):
                    chunks.append({
                        "content": chunk,
                        "source": os.path.basename(file_path),
                        "page": i + 1,
                        "doc_type": "xlsx"
                    })
        except Exception as e:
            logger.error(f"Error parsing XLSX {file_path}: {e}")
        return chunks

    def _parse_pptx(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            prs = Presentation(file_path)
            full_text = ""

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = f"\n\n--- Slide {slide_num} ---\n\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text += shape.text + "\n"
                full_text += slide_text

            text_chunks = self._split_into_chunks(full_text)
            for i, chunk in enumerate(text_chunks):
                chunks.append({
                    "content": chunk,
                    "source": os.path.basename(file_path),
                    "page": i + 1,
                    "doc_type": "pptx"
                })
        except Exception as e:
            logger.error(f"Error parsing PPTX {file_path}: {e}")
        return chunks

    def _parse_txt(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            with open(file_path, 'r', encoding=self.encoding, errors='replace') as f:
                content = f.read()

            text_chunks = self._split_into_chunks(content)
            for i, chunk in enumerate(text_chunks):
                chunks.append({
                    "content": chunk,
                    "source": os.path.basename(file_path),
                    "page": i + 1,
                    "doc_type": "txt"
                })
        except Exception as e:
            logger.error(f"Error parsing TXT {file_path}: {e}")
        return chunks

    def _parse_html(self, file_path: str) -> List[Dict[str, Any]]:
        chunks = []
        try:
            with open(file_path, 'r', encoding=self.encoding, errors='replace') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')
                text = soup.get_text(separator='\n', strip=True)

            text_chunks = self._split_into_chunks(text)
            for i, chunk in enumerate(text_chunks):
                chunks.append({
                    "content": chunk,
                    "source": os.path.basename(file_path),
                    "page": i + 1,
                    "doc_type": "html"
                })
        except Exception as e:
            logger.error(f"Error parsing HTML {file_path}: {e}")
        return chunks

    def _parse_md(self, file_path: str) -> List[Dict[str, Any]]:
        return self._parse_txt(file_path)

    def _split_into_chunks(self, text: str, chunk_size: Optional[int] = None, overlap: Optional[int] = None) -> List[str]:
        chunk_size = chunk_size or self.chunk_size
        overlap = overlap or self.chunk_overlap

        text = re.sub(r'[ \t]+', ' ', text).strip()
        if len(text) <= chunk_size:
            return [text] if text else []

        chunks = []
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]

            if end < len(text):
                last_period = chunk.rfind('。')
                last_newline = chunk.rfind('\n')
                split_pos = max(last_period, last_newline)
                if split_pos > chunk_size // 2:
                    chunk = chunk[:split_pos + 1]
                    end = start + split_pos + 1

            chunks.append(chunk.strip())
            start = end - overlap if end < len(text) else end

        return [c for c in chunks if c.strip()]


