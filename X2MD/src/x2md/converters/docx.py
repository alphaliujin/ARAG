from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn

from x2md.converters.base import BaseConverter
from x2md.llm import OllamaClient, llm_batch_strip_noise
from x2md.utils import (
    clean_markdown,
    extract_images_from_zip,
    get_image_output_dir,
    image_marker,
    split_contact_info,
    strip_toc_markers,
    table_to_lines,
)


class DocxConverter(BaseConverter):
    extensions = [".docx"]

    def convert(self, file_path: Path, **kwargs) -> str:
        llm_client: OllamaClient | None = kwargs.get("llm_client")
        md_output_path: Path | None = kwargs.get("md_output_path")

        image_dir = get_image_output_dir(md_output_path, file_path)
        image_map = extract_images_from_zip(
            file_path, "word/media/", image_dir, file_path.stem
        )

        try:
            doc = Document(str(file_path))
        except PackageNotFoundError as e:
            # 加密 docx / 已损坏的 zip
            raise RuntimeError(
                f"文件 {file_path.name} 无法解析(可能已加密或已损坏): {e}"
            ) from e
        header_footer_texts = self._collect_header_footer_texts(doc)

        paragraphs_data: list[dict] = []

        # 优化：直接遍历段落，不遍历XML元素
        for para in doc.paragraphs:
            text = para.text.strip()

            if self._is_toc_field(para):
                continue

            if text in header_footer_texts:
                continue

            style_name = (para.style.name or "").lower() if para.style else ""
            paragraphs_data.append({
                "text": text,
                "style": style_name,
                "para": para,
            })

        # 优化：直接使用doc.tables，不遍历XML元素
        tables_data: list[dict] = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            if rows:
                tables_data.append({"rows": rows})

        noise_indices: set[int] = set()
        if llm_client and llm_client.is_available():
            candidate_texts = []
            candidate_indices = []
            for i, pd in enumerate(paragraphs_data):
                text = pd["text"]
                if not text:
                    continue
                style = pd["style"]
                if "heading" in style:
                    continue
                candidate_texts.append(text)
                candidate_indices.append(i)

            if candidate_texts:
                batch_noise = llm_batch_strip_noise(llm_client, candidate_texts)
                for local_idx in batch_noise:
                    noise_indices.add(candidate_indices[local_idx])

        parts: list[str] = []
        for i, pd in enumerate(paragraphs_data):
            text = pd["text"]
            style = pd["style"]
            para = pd["para"]

            if i in noise_indices:
                continue

            # 收集本段内的图片标记 (不直接 append 到 parts,避免顺序倒置)
            drawing_markers = self._collect_drawing_markers(para, image_map)

            if not text:
                # 空段落但带图片 → 仅追加图片
                if drawing_markers:
                    parts.extend(drawing_markers)
                else:
                    parts.append("")
                continue

            # 先追加文本(带 heading/list 前缀),再追加该段对应的图片
            # 这样得到 [文本, 图片1, 图片2] 的阅读顺序;
            # 旧实现 _check_drawing 在文本 append 之前已先 append 图片,顺序错乱。
            if "heading 1" in style:
                parts.append(f"# {text}")
            elif "heading 2" in style:
                parts.append(f"## {text}")
            elif "heading 3" in style:
                parts.append(f"### {text}")
            elif "heading 4" in style:
                parts.append(f"#### {text}")
            elif "heading 5" in style:
                parts.append(f"##### {text}")
            elif "heading 6" in style:
                parts.append(f"###### {text}")
            elif "list" in style:
                parts.append(f"- {text}")
            else:
                parts.append(text)

            if drawing_markers:
                parts.extend(drawing_markers)

        for td in tables_data:
            md_table = table_to_lines(td["rows"])
            if md_table:
                parts.append(md_table)

        result = "\n\n".join(parts)
        result = strip_toc_markers(result)
        result = split_contact_info(result)
        return clean_markdown(result)

    @staticmethod
    def _collect_header_footer_texts(doc: Document) -> set[str]:
        texts: set[str] = set()
        for section in doc.sections:
            for para in section.header.paragraphs:
                t = para.text.strip()
                if t:
                    texts.add(t)
            for para in section.footer.paragraphs:
                t = para.text.strip()
                if t:
                    texts.add(t)
            if section.first_page_header:
                for para in section.first_page_header.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.add(t)
            if section.first_page_footer:
                for para in section.first_page_footer.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.add(t)
        return texts

    @staticmethod
    def _is_toc_field(para) -> bool:
        # 优化：只查直接子 w:r 下的 w:instrText，不再 iter() 整棵子树
        for r in para._element.findall(qn("w:r")):
            for instr in r.findall(qn("w:instrText")):
                if "TOC" in (instr.text or ""):
                    return True
        return False

    @staticmethod
    def _collect_drawing_markers(para, image_map: dict[str, str]) -> list[str]:
        """收集段落中的图片标记,按出现顺序返回,不直接修改外部 parts."""
        markers: list[str] = []
        for child in para._element.iter():
            if child.tag == qn("a:blip"):
                embed = child.get(qn("r:embed"))
                if embed:
                    rel = para.part.rels.get(embed)
                    if rel and hasattr(rel, "target_ref"):
                        target = rel.target_ref
                        if target in image_map:
                            markers.append(image_marker(image_map[target]))
                        elif target.startswith("media/"):
                            media_name = f"word/{target}"
                            if media_name in image_map:
                                markers.append(image_marker(image_map[media_name]))
        return markers

