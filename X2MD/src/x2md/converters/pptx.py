from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from x2md.converters.base import BaseConverter
from x2md.utils import (
    clean_markdown,
    get_image_output_dir,
    image_marker,
    split_contact_info,
    table_to_lines,
)


class PptxConverter(BaseConverter):
    extensions = [".pptx"]

    def convert(self, file_path: Path, **kwargs) -> str:
        md_output_path: Path | None = kwargs.get("md_output_path")
        image_dir = get_image_output_dir(md_output_path, file_path)
        image_dir.mkdir(parents=True, exist_ok=True)

        prs = Presentation(str(file_path))
        parts: list[str] = []
        image_counter = 0

        for i, slide in enumerate(prs.slides):
            slide_parts: list[str] = []
            slide_parts.append(f"## Slide {i + 1}")

            for shape in slide.shapes:
                if shape.shape_type == 13:
                    image_counter += 1
                    image_name = self._save_shape_image(
                        shape, image_dir, file_path.stem, image_counter
                    )
                    if image_name:
                        slide_parts.append(image_marker(image_name))
                    continue

                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if not text:
                            continue
                        level = para.level
                        if level == 0:
                            slide_parts.append(text)
                        else:
                            indent = "  " * level
                            slide_parts.append(f"{indent}- {text}")

                if shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        rows.append([cell.text.strip() for cell in row.cells])
                    md_table = table_to_lines(rows)
                    if md_table:
                        slide_parts.append(md_table)

            parts.append("\n\n".join(slide_parts))

        result = "\n\n---\n\n".join(parts)
        result = split_contact_info(result)
        return clean_markdown(result)

    @staticmethod
    def _save_shape_image(
        shape, image_dir: Path, stem: str, counter: int
    ) -> str | None:
        try:
            img = shape.image
            content_type = img.content_type
            blob = img.blob
            ext_map = {
                "image/png": ".png",
                "image/jpeg": ".jpg",
                "image/gif": ".gif",
                "image/bmp": ".bmp",
                "image/tiff": ".tiff",
                "image/webp": ".webp",
            }
            # 旧实现保留 emf/wmf 作为合法扩展名,但 PIL 在 macOS/Linux 无法打开 EMF/WMF,
            # 下游 ViT (Image.open(...).convert("RGB")) 会抛异常并让整批失败。
            # 直接跳过这两种格式,避免落磁盘后造成"images 目录里有死文件"。
            if content_type in ("image/emf", "image/wmf", "image/x-emf", "image/x-wmf"):
                return None
            ext = ext_map.get(content_type, ".png")
            filename = f"{stem}_{counter:03d}{ext}"
            filepath = image_dir / filename
            filepath.write_bytes(blob)
            return filename
        except Exception:
            return None
