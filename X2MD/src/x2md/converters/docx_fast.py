"""快速DOCX转换器 - 优化性能版本"""

from __future__ import annotations
from pathlib import Path
from docx import Document
from x2md.converters.base import BaseConverter
from x2md.utils import clean_markdown


class DocxFastConverter(BaseConverter):
    """高性能DOCX转换器，跳过图片提取和LLM处理"""

    extensions = [".docx"]

    def convert(self, file_path: Path, **kwargs) -> str:
        """快速转换DOCX为Markdown"""
        doc = Document(str(file_path))

        parts: list[str] = []

        # 直接遍历段落，不处理图片
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # 简单样式判断
            style_name = ""
            if para.style and para.style.name:
                style_name = para.style.name.lower()

            if "heading 1" in style_name:
                parts.append(f"# {text}")
            elif "heading 2" in style_name:
                parts.append(f"## {text}")
            elif "heading 3" in style_name:
                parts.append(f"### {text}")
            elif "list" in style_name:
                parts.append(f"- {text}")
            else:
                parts.append(text)

        # 简单表格处理
        for table in doc.tables:
            rows = []
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells]
                if any(row_texts):
                    rows.append("| " + " | ".join(row_texts) + " |")

            if len(rows) >= 2:
                # 添加表头分隔符
                col_count = len(table.rows[0].cells)
                separator = "|" + "---|" * col_count
                rows.insert(1, separator)
                parts.extend(rows)

        result = "\n\n".join(parts)
        return clean_markdown(result)
