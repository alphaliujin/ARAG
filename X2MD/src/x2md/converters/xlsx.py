from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from x2md.converters.base import BaseConverter
from x2md.utils import clean_markdown, split_contact_info, table_to_lines


class XlsxConverter(BaseConverter):
    extensions = [".xlsx"]

    def convert(self, file_path: Path, **kwargs) -> str:
        include_all_sheets: bool = kwargs.get("include_all_sheets", True)
        sheet_name: str | None = kwargs.get("sheet_name", None)

        try:
            wb = load_workbook(str(file_path), read_only=True, data_only=True)
        except InvalidFileException as e:
            raise RuntimeError(
                f"文件 {file_path.name} 无法解析(可能已加密或文件格式不正确): {e}"
            ) from e
        except Exception as e:
            # openpyxl 对加密 xlsx 偶尔抛 zipfile.BadZipFile,统一友好提示
            msg = str(e).lower()
            if "encrypted" in msg or "password" in msg or "bad zip" in msg:
                raise RuntimeError(
                    f"文件 {file_path.name} 似乎已加密或损坏: {e}"
                ) from e
            raise
        try:
            parts: list[str] = []

            target_sheets = [sheet_name] if sheet_name else wb.sheetnames
            if not include_all_sheets and not sheet_name:
                target_sheets = target_sheets[:1]

            for name in target_sheets:
                if name not in wb.sheetnames:
                    continue
                ws = wb[name]

                if len(target_sheets) > 1:
                    parts.append(f"## Sheet: {name}")

                # 注: read_only 模式不支持 merged_cells.ranges 的精确展开,
                # 这里只能保留 cell 原值;非 read_only 模式才能合并展开,但内存开销大。
                # 折中: 文档转换以"信息可读"为优先,合并单元格只在表头第一行常见,
                # 当前实现把空 cell 留空,下游 RAG 阅读时仍可推断含义。
                rows: list[list[str]] = []
                for row in ws.iter_rows(values_only=True):
                    rows.append([str(cell) if cell is not None else "" for cell in row])

                if rows:
                    md_table = table_to_lines(rows)
                    if md_table:
                        parts.append(md_table)

            result = "\n\n".join(parts)
            result = split_contact_info(result)
            return clean_markdown(result)
        finally:
            wb.close()
